"""Construit l'Excel à envoyer à Ndèye Maguette, à partir de l'atlas officiel importé et de reponses.csv.

Onglets : Guide · Phrases de Kaddu (pulaar et sérère à remplir) · Questions · Doublons du lexique ·
Montants à vérifier. Les cellules jaunes sont celles à remplir.

Usage (depuis la racine du dépôt) : python nlp/outils/excel_pour_ndeye.py [sortie.xlsx]
"""
import csv
import sys
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

ATLAS = Path(__file__).resolve().parent.parent / "data" / "atlas"
SORTIE = Path(sys.argv[1]) if len(sys.argv) > 1 else Path("TEWLU_Kaddu_pour_Ndeye.xlsx")

POLICE = "Arial"
TITRE = Font(name=POLICE, bold=True, size=14, color="1E3A8A")
ENTETE = Font(name=POLICE, bold=True, color="FFFFFF")
NORMAL = Font(name=POLICE, size=10)
GRAS = Font(name=POLICE, size=10, bold=True)
FOND_ENTETE = PatternFill("solid", start_color="3B82F6")
JAUNE = PatternFill("solid", start_color="FFF59D")
GRIS = PatternFill("solid", start_color="F1F5F9")
BORD = Border(*(Side(style="thin", color="CBD5E1"),) * 4)
HAUT = Alignment(wrap_text=True, vertical="top")


def lire(nom):
    with open(ATLAS / nom, encoding="utf-8") as f:
        return list(csv.DictReader(f))


def tableau(ws, ligne, entetes, lignes, largeurs, a_remplir=()):
    for j, (h, l) in enumerate(zip(entetes, largeurs), start=1):
        c = ws.cell(ligne, j, h)
        c.font, c.fill, c.alignment, c.border = ENTETE, FOND_ENTETE, HAUT, BORD
        ws.column_dimensions[c.column_letter].width = l
    for i, valeurs in enumerate(lignes, start=ligne + 1):
        for j, v in enumerate(valeurs, start=1):
            c = ws.cell(i, j, v)
            c.font, c.alignment, c.border = NORMAL, HAUT, BORD
            if j in a_remplir:
                c.fill = JAUNE
    ws.freeze_panes = ws.cell(ligne + 1, 1)


wb = Workbook()

# --- Guide
ws = wb.active
ws.title = "Guide"
ws.column_dimensions["A"].width = 110
lignes_guide = [
    ("TEWLU Kaddu — ce qu'il reste à valider", TITRE),
    ("Bonjour Ndèye ! L'atlas officiel que tu as validé (Off_Annuaire_..._3.xlsx) est maintenant le seul utilisé "
     "par Kaddu. Ce fichier rassemble ce qui manque encore.", NORMAL),
    ("", NORMAL),
    ("Quoi remplir : uniquement les cellules JAUNES.", GRAS),
    ("1. « Phrases de Kaddu » : ce que Kaddu répond. Pour chaque ligne, écris la phrase en pulaar et en sérère "
     "dans les colonnes jaunes (ou corrige la version actuelle, reprise de ton document « Phases Tests »).", NORMAL),
    ("   Garde les mots entre accolades tels quels : {client}, {montant}, {articles}, {reste}, {ventes}, {dettes}. "
     "Kaddu les remplace (exemple : {client} -> Hawa, {montant} -> trois mille francs).", NORMAL),
    ("   Exemple : « {client} yobi {montant}, hedo {reste}. Ey walla ala ? » devient « Hawa yobi mille francs, "
     "hedo deux mille francs. Ey walla ala ? ».", NORMAL),
    ("   Une phrase laissée vide est dite en français par Kaddu (signalée en orange dans l'espace de test).", NORMAL),
    ("2. « Questions » : des points à trancher. Réponds dans la colonne jaune.", NORMAL),
    ("3. « Doublons du lexique » : un même mot rattaché à deux concepts dans la même langue. Dis lequel garder.",
     NORMAL),
    ("4. « Montants à vérifier » : des variantes de montants qui sont en fait la forme d'un autre montant "
     "(souvent un début de forme coupé). Donne la bonne variante ou écris « supprimer ».", NORMAL),
    ("", NORMAL),
    ("Pour tester Kaddu toi-même : ouvre le lien de l'espace de test (Marème te l'envoie), choisis la langue, "
     "parle ou écris, puis réponds oui / non.", NORMAL),
]
for i, (texte, police) in enumerate(lignes_guide, start=1):
    c = ws.cell(i, 1, texte)
    c.font, c.alignment = police, Alignment(wrap_text=True, vertical="top")
c = ws.cell(len(lignes_guide) + 2, 1, "Cellule jaune = à remplir")
c.fill, c.font = JAUNE, GRAS

# --- Phrases de Kaddu
ws = wb.create_sheet("Phrases de Kaddu")
ws.cell(1, 1, "Phrases de Kaddu — pulaar et sérère à remplir (colonnes jaunes)").font = TITRE
reponses = lire("reponses.csv")
tableau(ws, 3, ["clé", "quand Kaddu le dit", "mots à garder", "français", "wolof (actuel)", "pulaar actuel",
                "PULAAR (à remplir / corriger)", "sérère actuel", "SÉRÈRE (à remplir / corriger)", "commentaire"],
        [[r["cle"], r["usage"], r["variables"].replace(";", ", "), r["fr"], r["wo"], r["ff"], r["ff"], r["srr"],
          r["srr"], ""] for r in reponses],
        [24, 30, 16, 40, 40, 30, 34, 30, 34, 30], a_remplir=(7, 9, 10))

# --- Questions
ws = wb.create_sheet("Questions")
ws.cell(1, 1, "Questions à trancher").font = TITRE
questions = [
    ("Montants", "Quel mot dire entre le montant en français et le montant en langue (« maanaam » en wolof) ? "
     "Pulaar : ? Sérère : ?", "Tant que c'est vide, Kaddu ne dit que le montant en français en pulaar et en sérère."),
    ("Montants", "Pulaar, 150 000 F : l'atlas dit « thiapantati ujunéré », mais la règle des autres lignes donnerait "
     "« ujunaajé thiapanɗe tati ». Quelle forme ? Et pour 200 000 F, 250 000 F... ?", "Onglet Pulaar_Montants."),
    ("Montants", "Sérère, 100 000 F : « thiouni xarbéndik » (atlas) ou « thiouni xarbéne dik » ? Les deux ?", ""),
    ("Phases Tests", "pu_002 pulaar : « téemédé tati » = 1 500 F d'après l'atlas, mais la traduction dit 3 000 F et la "
     "réponse attendue dit « teémédeé jeegom » (3 000 F). Quelle est la bonne phrase ?", ""),
    ("Phases Tests", "Sérère : le document écrit « teeméte » (avec t), l'atlas « teeméde » (avec d). Faut-il ajouter "
     "« teeméte » comme variante ?", "Kaddu reconnaît aujourd'hui seulement « teeméde »."),
    ("Phases Tests", "Sérère pu_004 : que veut dire la fin « ngess naa boo ndiik » ? Faut-il la garder dans la "
     "réponse de Kaddu ?", "Pour l'instant, Kaddu dit seulement « Hawa a sakha nang {montant}. »"),
    ("Phases Tests", "Pulaar pu_005 : la réponse donne seulement les créances (« Ada jogui ... »). Faut-il aussi "
     "dire les ventes du jour, comme en wolof et en sérère ?", ""),
    ("Phrases", "Quantités : Kaddu dit les quantités en français (« dix saak maalo »). Faut-il les dire dans la "
     "langue (« xarbaxay saak maalo ») ?", ""),
    ("Phrases", "Sérère : quel mot pour relier deux produits (« riz et sucre ») ?", "Pulaar : « e » (à confirmer)."),
    ("Phrases", "Message avec deux opérations (une ancienne dette et une nouvelle vente) : comment le commerçant "
     "le dit-il ? Exemples dans chaque langue ?", ""),
]
tableau(ws, 3, ["n°", "thème", "question", "note", "RÉPONSE DE NDÈYE"],
        [[i, t, q, n, ""] for i, (t, q, n) in enumerate(questions, start=1)], [5, 14, 70, 45, 50], a_remplir=(5,))

# --- Doublons du lexique
ws = wb.create_sheet("Doublons du lexique")
ws.cell(1, 1, "Même mot pour deux concepts dans la même langue").font = TITRE
doublons = lire("doublons_atlas.csv")
noms = {"wo": "wolof", "ff": "pulaar", "srr": "sérère"}
tableau(ws, 3, ["langue", "mot (écrit simplifié)", "concepts", "QUEL CONCEPT GARDER ?"],
        [[noms[d["langue"]], d["mot"], d["concepts"], ""] for d in doublons], [10, 22, 45, 45], a_remplir=(4,))

# --- Montants à vérifier
ws = wb.create_sheet("Montants à vérifier")
ws.cell(1, 1, "Variantes de montants à vérifier").font = TITRE
montants = lire("montants_a_verifier.csv")
tableau(ws, 3, ["langue", "montant (F)", "forme de l'atlas", "variante de l'atlas", "problème",
                "BONNE VARIANTE (ou « supprimer »)"],
        [[noms[m["langue"]], int(m["montant_fcfa"]), m["forme"], m["variante"], m["probleme"], ""] for m in montants],
        [10, 12, 36, 26, 24, 40], a_remplir=(6,))

for feuille in wb.worksheets:
    feuille.sheet_view.showGridLines = feuille.title != "Guide"
wb.save(SORTIE)
print("Écrit :", SORTIE)
